#!/usr/bin/env python3
"""Generate the reviewer-corrected manuscript, supplement, PDF, and editable DOCX."""
from __future__ import annotations

import json
import re
import shutil
import subprocess
from pathlib import Path

import numpy as np
import pandas as pd
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches


ROOT = Path(__file__).resolve().parents[1]
PAPER = ROOT / "paper"


def value(row: dict, key: str, digits: int = 2) -> str:
    return f"{float(row[key]):.{digits}f}"


def replace(text: str, values: dict[str, object]) -> str:
    for key, item in values.items():
        text = text.replace(f"@@{key}@@", str(item))
    if "@@" in text:
        unresolved = sorted({part.split("@@", 1)[0] for part in text.split("@@")[1::2]})
        raise RuntimeError(f"Unresolved manuscript tokens: {unresolved}")
    return text


def add_docx_table(document: Document, frame: pd.DataFrame, columns: list[str]) -> None:
    table = document.add_table(rows=1, cols=len(columns)); table.style = "Table Grid"
    for i, column in enumerate(columns): table.rows[0].cells[i].text = column
    for row in frame[columns].itertuples(index=False):
        cells = table.add_row().cells
        for i, item in enumerate(row): cells[i].text = f"{item:.3f}" if isinstance(item, float) else str(item)


def tex_plain(text: str) -> str:
    text = re.sub(r"\\cite\{([^}]*)\}", lambda m: "[" + m.group(1).replace(",", ", ") + "]", text)
    replacements = {r"\%": "%", r"\_": "_", r"\geq": "≥", r"\leq": "≤", r"\times": "×",
                    r"\pm": "±", r"\alpha": "alpha", r"\textwidth": "text width", "--": "–"}
    for old, new in replacements.items(): text = text.replace(old, new)
    # Preserve the content of common formatting commands.
    for _ in range(4):
        text = re.sub(r"\\(?:emph|texttt|textbf|mathrm|mathbf)\{([^{}]*)\}", r"\1", text)
    text = re.sub(r"\\[A-Za-z]+\*?(?:\[[^]]*\])?", "", text)
    text = text.replace("{", "").replace("}", "").replace("$", "")
    text = re.sub(r"\s+", " ", text).strip()
    return text


def add_complete_tex_narrative(document: Document, tex: str) -> None:
    abstract = re.search(r"\\begin\{abstract\}(.*?)\\end\{abstract\}", tex, flags=re.S)
    document.add_heading("Abstract", level=1)
    document.add_paragraph(tex_plain(abstract.group(1)) if abstract else "")
    start = tex.index(r"\section{Introduction}")
    stop = tex.index(r"\begin{thebibliography}")
    body = tex[start:stop]
    body = re.sub(r"\\begin\{figure\*?\}.*?\\end\{figure\*?\}", "", body, flags=re.S)
    body = re.sub(r"\\begin\{table\*?\}.*?\\end\{table\*?\}", "", body, flags=re.S)
    headings = list(re.finditer(r"\\(section|subsection)\{([^}]*)\}", body))
    for i, match in enumerate(headings):
        document.add_heading(tex_plain(match.group(2)), level=1 if match.group(1) == "section" else 2)
        end = headings[i + 1].start() if i + 1 < len(headings) else len(body)
        content = body[match.end():end]
        for paragraph in re.split(r"\n\s*\n", content):
            cleaned = tex_plain(paragraph)
            if cleaned: document.add_paragraph(cleaned)
    references = tex[stop:tex.index(r"\end{thebibliography}")]
    document.add_heading("References", level=1)
    entries = re.split(r"\\bibitem\{[^}]*\}", references)[1:]
    for number, entry in enumerate(entries, 1):
        cleaned = tex_plain(entry)
        if cleaned: document.add_paragraph(f"[{number}] {cleaned}")


def main() -> None:
    PAPER.mkdir(parents=True, exist_ok=True)
    shutil.copy2(ROOT / "templates/IEEEtran.cls", PAPER / "IEEEtran.cls")
    nested = json.loads((ROOT / "phase2/outputs/nested200/DECISION.json").read_text())
    selected = nested["summary"]["nested_selected_procedure"]
    prior_nested = nested["summary"]["historical_fixed_prior"]
    mean_nested = nested["summary"]["outer_fold_training_mean"]
    repeated_decision = json.loads((ROOT / "phase2/outputs/innovation200_confirmation/DECISION.json").read_text())
    repeated = repeated_decision["winner_oof"]
    vital = json.loads((ROOT / "phase2/outputs/innovation200_vitaldb_external/DECISION.json").read_text())
    event = nested["event_safety"]
    conformal = pd.read_csv(ROOT / "phase2/outputs/icbme_clinical_analysis/CONFORMAL_90.csv")
    split = json.loads((ROOT / "data/mimic_bp/splits/official_patient_split.json").read_text())
    with np.load(ROOT / "outputs/neural/cache/metadata.npz", allow_pickle=False) as metadata:
        segment_counts = {
            name: int(np.sum(metadata["split"].astype(str) == name))
            for name in ("train", "validation", "test")
        }
    selected_names = ", ".join(item["candidate_id"].replace("_", r"\_") for item in nested["selection"])
    high = event["high_bp"]; low = event["map_below_65"]
    conf = conformal.iloc[0]
    paired = nested["paired_patient_tests"]["selected_vs_fold_mean"]
    values = {
        "NEST_SBP": value(selected, "mae_sbp"), "NEST_DBP": value(selected, "mae_dbp"),
        "NEST_RSBP": value(selected, "rmse_sbp"), "NEST_RDBP": value(selected, "rmse_dbp"),
        "NEST_BSBP": value(selected, "bias_sbp"), "NEST_BDBP": value(selected, "bias_dbp"),
        "NEST_PSBP": value(selected, "pearson_sbp"), "NEST_PDBP": value(selected, "pearson_dbp"),
        "PRIOR_SBP": value(prior_nested, "mae_sbp"), "PRIOR_DBP": value(prior_nested, "mae_dbp"),
        "MEAN_SBP": value(mean_nested, "mae_sbp"), "MEAN_DBP": value(mean_nested, "mae_dbp"),
        "DELTA": f"{paired['patient_mean_mae_delta_candidate_minus_baseline']:.2f}",
        "DELTA_LO": f"{paired['ci_low']:.2f}", "DELTA_HI": f"{paired['ci_high']:.2f}", "DELTA_P": f"{paired['wilcoxon_p']:.3g}",
        "REP_SBP": value(repeated, "mae_sbp"), "REP_DBP": value(repeated, "mae_dbp"),
        "HIGH_PREV": f"{100*high['direct']['prevalence']:.1f}", "HIGH_SENS": f"{100*high['direct']['sensitivity']:.1f}",
        "HIGH_SPEC": f"{100*high['direct']['specificity']:.1f}", "HIGH_AUPRC": f"{high['auprc']:.3f}", "HIGH_AUROC": f"{high['auroc']:.3f}",
        "HIGH_XSENS": f"{100*high['crossfit_90pct_specificity']['sensitivity']:.1f}",
        "LOW_PREV": f"{100*low['direct']['prevalence']:.1f}", "LOW_SENS": f"{100*low['direct']['sensitivity']:.1f}",
        "LOW_SPEC": f"{100*low['direct']['specificity']:.1f}", "LOW_AUPRC": f"{low['auprc']:.3f}", "LOW_AUROC": f"{low['auroc']:.3f}",
        "LOW_XSENS": f"{100*low['crossfit_90pct_specificity']['sensitivity']:.1f}",
        "VZ_SBP": value(vital["zero_shot"], "mae_sbp"), "VZ_DBP": value(vital["zero_shot"], "mae_dbp"),
        "VZ_BSBP": value(vital["zero_shot"], "bias_sbp"), "VZ_BDBP": value(vital["zero_shot"], "bias_dbp"),
        "VM_SBP": value(vital["source_mean"], "mae_sbp"), "VM_DBP": value(vital["source_mean"], "mae_dbp"),
        "VA_SBP": value(vital["supervised_affine_context"], "mae_sbp"), "VA_DBP": value(vital["supervised_affine_context"], "mae_dbp"),
        "V_AUPRC": f"{vital['zero_shot']['auprc']:.3f}", "V_AUROC": f"{vital['zero_shot']['auroc']:.3f}",
        "V_XSENS": f"{100*vital['crossfit_event_threshold']['sensitivity']:.1f}", "V_XSPEC": f"{100*vital['crossfit_event_threshold']['specificity']:.1f}",
        "CONF_SW": f"{float(conf['sbp_mean_width']):.2f}", "CONF_DW": f"{float(conf['dbp_mean_width']):.2f}",
        "CONF_SC": f"{100*float(conf['sbp_coverage']):.1f}", "CONF_DC": f"{100*float(conf['dbp_coverage']):.1f}",
        "SELECTED_NAMES": selected_names,
        "TRAIN_SEG": segment_counts["train"], "VAL_SEG": segment_counts["validation"],
        "TEST_SEG": segment_counts["test"],
    }
    zero_mean_mae = (float(vital["zero_shot"]["mae_sbp"]) + float(vital["zero_shot"]["mae_dbp"])) / 2
    baseline_mean_mae = (float(vital["source_mean"]["mae_sbp"]) + float(vital["source_mean"]["mae_dbp"])) / 2
    values["V_RELATION"] = "worse than" if zero_mean_mae >= baseline_mean_mae else "better than"
    values["V_CONCLUSION"] = "failed to beat" if zero_mean_mae >= baseline_mean_mae else "outperformed"

    template = r'''\documentclass[conference]{IEEEtran}
\IEEEoverridecommandlockouts
\usepackage{cite,amsmath,amssymb,graphicx,booktabs,url}
\begin{document}
\title{Patient-Nested Selection and Cross-Domain Safety Audit of PPG-Only Blood-Pressure Estimation}
\author{\IEEEauthorblockN{Anonymous submission}}
\maketitle
\begin{abstract}
Photoplethysmography (PPG)-only blood-pressure estimation can obtain moderate aggregate error while missing clinically important pressure extremes. We trained a component-wise model family and evaluated the entire locked 200-configuration selection procedure using five patient-nested outer folds on 1,524 MIMIC-BP subjects (45,689 duplicate-audited segments). The nested selected procedure achieved systolic/diastolic (SBP/DBP) MAE @@NEST_SBP@@/@@NEST_DBP@@ mmHg versus @@MEAN_SBP@@/@@MEAN_DBP@@ for an outer-training-fold mean. A separate, explicitly descriptive 75-fit repeated grouped audit gave @@REP_SBP@@/@@REP_DBP@@ mmHg, yet direct high-BP sensitivity was only @@HIGH_SENS@@\% at @@HIGH_PREV@@\% prevalence. The fixed prior retained by that secondary audit, rather than the fold-varying nested procedure, was evaluated on a deterministic synchronized VitalDB stress test (119 patients; 2,380 windows). Its zero-shot MAE was @@VZ_SBP@@/@@VZ_DBP@@, @@V_RELATION@@ @@VM_SBP@@/@@VM_DBP@@ for the official source-training mean; supervised patient-cross-fitted target-domain adaptation reached @@VA_SBP@@/@@VA_DBP@@. These experiments show that patient-nested selection, range-specific error, simple baselines, and external calibration status materially change the interpretation of PPG-based BP accuracy. This is a retrospective estimation study, not validation of a diagnostic device.
\end{abstract}
\begin{IEEEkeywords}blood pressure estimation, clinical safety, domain shift, nested cross-validation, photoplethysmography\end{IEEEkeywords}

\section{Introduction}
PPG offers convenient peripheral pulse measurement, but absolute BP is only indirectly identifiable from pulse morphology and is affected by vascular tone, contact, sensor response, treatment, and population. Segment-wise evaluation can also expose subject identity or overlapping signal content; it does not follow that every published study leaks, but the risk makes patient-level separation essential \cite{gonzalez2023,cisnal2026}. Cross-dataset benchmarks report large calibration-free degradation even for strong one-dimensional networks \cite{moulaeifard2025}.

We asked a practical biomedical question: after a broad, reproducible architecture/preprocessing/objective search, does a trained PPG-only model improve over simple baselines under patient-nested evaluation, preserve performance at clinically important ranges, and transfer without target labels? Unlike a review, this work trains and evaluates models. Its novelty is the joint audit of a locked 200-method selection process, exact patient/duplicate controls, operational event performance, and synchronized external transfer. Negative results are retained rather than replaced by a post-hoc winner.

\section{Materials and Methods}
\subsection{Cohorts, targets, and quality control}
MIMIC-BP contains 1,524 ICU subjects with thirty 30-s, 125-Hz PPG segments per subject and published segment-level SBP/DBP labels derived from the paired arterial-pressure source \cite{sanches2024}. We pair each published 30-s label with the central 15 s of that same PPG segment; the local release does not contain raw ABP, so labels are not recomputed for the shorter input. Exact equality of all 3,750 raw PPG samples (BLAKE2b hash) revealed 30 duplicate groups. Both members of the one cross-patient group were removed; within-patient groups retained one canonical copy, leaving 45,689 segments.

The official subject-disjoint split is 1,100/195/229 train/validation/test, corresponding after duplicate control to @@TRAIN_SEG@@/@@VAL_SEG@@/@@TEST_SEG@@ segments. It was used for development screening. The primary nested analysis subsequently repartitioned all 1,524 subjects; thus formerly designated test subjects are included in outer OOF evaluation and are not called untouched. Split ID lists and exclusions are released.

External tests comprised PPG-BP (219 people; short PPG with cuff labels), an ambulatory cuff-labelled cohort (56 people), and VitalDB \cite{liang2018,vasquez2024,arguello2025,lee2022}. Cuff labels are not simultaneous invasive references. VitalDB cases were ranked outcome-blind by SHA-256 over a fixed string; 128 cases with PLETH, ART\_SBP, and ART\_DBP tracks were selected. Timestamp-aligned 500-Hz PLETH was resampled to 125 Hz, and each non-overlapping 15-s window used the medians of contemporaneous numeric ART\_SBP/ART\_DBP tracks. Fixed QC accepted 119 patients and 2,380 windows (maximum 20 per case). This changes device, operating-room population, care processes, and signal pipeline together and therefore measures combined domain shift, not isolated device shift.

\subsection{Inputs, models, and locked search}
Each input was robustly scaled by its own median/IQR; VPG and APG were numerical derivatives. Waveform channels and every engineered feature were computed from the same central 15-s source window. Candidate views included raw/normalized PPG, PPG/VPG/APG, smoothed and detrended channels, 2--15-s duration, and engineered morphology. Backbones included FCN, residual, temporal/gated convolution, Inception-like, ConvNeXt-1D, CNN--BiGRU, patch Transformer/mixer, long-kernel, morphology-stem, multiscale, and temporal--spectral models. PaPaGei-S was tested as a PPG-native foundation encoder \cite{pillai2025}; a text LLM was excluded because neither its tokenizer nor pretraining objective matches continuous waveforms.

The direct output regressed SBP/DBP. Auxiliary variants predicted pulse pressure and the approximation $MAP_a=(SBP+2DBP)/3$; $MAP_a$ is not presented as measured mean arterial pressure. Objectives included Huber/L1/log-cosh, concordance, asymmetric high-BP, focal, ordinal/clinical, tail weighting, and physiological multitask terms; training varied optimizers, schedules, sampling, and perturbations. The 200 configurations (80 architecture, 40 input, 40 objective, 40 training) were literature-informed and systematically constructed. Component summaries are associations across a non-factorial search, not causal ablations. Every configuration and result is supplied.

\subsection{Patient-nested evaluation and secondary confirmation}
The primary estimate used five stratified patient-level outer folds. Within each outer-training set, a fixed-seed 80/20 patient split screened all 200 locked configurations plus the historical prior (1,005 inner fits). The lowest normalized validation score set architecture and epoch count; that configuration was retrained on the full outer-training set for the fixed number of epochs and evaluated once on the outer test subjects. No patient crossed inner/outer boundaries and every segment received exactly one outer prediction. Because the configuration registry itself followed prior interaction with this cohort, even this estimates a locked procedure on the current population rather than a never-seen prospective cohort.

For continuity with earlier experiments, four preselected finalists and a frozen prior also underwent five patient-wise folds with three new seeds (75 fits). Since design selection preceded these folds, this repeated OOF result is explicitly secondary/descriptive, not a fully independent estimate. Patient-macro comparisons used 20,000 clustered bootstrap resamples and paired Wilcoxon tests. Population means were fitted only on each training fold.

\subsection{Safety, external transfer, and uncertainty}
High BP was an operational endpoint, SBP$\geq140$ or DBP$\geq90$ mmHg, not a universal diagnostic definition. Low derived pressure was $MAP_a<65$ mmHg and is called a low-MAP event, not low perfusion pressure. We report sensitivity, specificity, AUROC/AUPRC, pressure-stratified MAE/bias, mean error and SD, 95\% Bland--Altman limits, and error-within-5/10/15-mmHg rates. External ``zero-shot'' uses no destination BP labels. The external checkpoint ensemble is the fixed prior retained by the secondary confirmation and is not the fold-varying nested-selection procedure. The source baseline is calculated from the official internal training partition only. Affine ridge output adaptation is supervised target-domain adaptation fitted by patient-level cross-fitting; it is neither zero-shot nor subject-specific calibration. Cross-fitted residual quantiles produced an exploratory 90\% interval audit. Ensemble dispersion is only a heuristic abstention score, not calibrated uncertainty.

\section{Results}
\subsection{Nested model selection}
The five inner folds selected: @@SELECTED_NAMES@@. The resulting patient-nested OOF MAE was @@NEST_SBP@@/@@NEST_DBP@@ mmHg, RMSE @@NEST_RSBP@@/@@NEST_RDBP@@, bias @@NEST_BSBP@@/@@NEST_BDBP@@, and correlation @@NEST_PSBP@@/@@NEST_PDBP@@. The historical fixed prior gave @@PRIOR_SBP@@/@@PRIOR_DBP@@, while outer-training means gave @@MEAN_SBP@@/@@MEAN_DBP@@ (Table~\ref{tab:nested}). The selected procedure minus mean patient-level MAE difference was @@DELTA@@ mmHg (95\% CI @@DELTA_LO@@ to @@DELTA_HI@@; Wilcoxon $p=$@@DELTA_P@@). The secondary repeated grouped estimate for the retained prior was @@REP_SBP@@/@@REP_DBP@@.

\begin{table}[t]\caption{Primary patient-nested OOF performance (mmHg)}\label{tab:nested}\centering\scriptsize
\begin{tabular}{lrrrr}\toprule Model & SBP MAE & DBP MAE & SBP RMSE & DBP RMSE\\\midrule
Nested selected procedure & @@NEST_SBP@@ & @@NEST_DBP@@ & @@NEST_RSBP@@ & @@NEST_RDBP@@\\
Historical fixed prior & @@PRIOR_SBP@@ & @@PRIOR_DBP@@ & -- & --\\
Outer-training mean & @@MEAN_SBP@@ & @@MEAN_DBP@@ & -- & --\\\bottomrule\end{tabular}\end{table}

\subsection{Clinical safety}
On the primary nested OOF predictions, direct high-BP sensitivity was @@HIGH_SENS@@\%, specificity @@HIGH_SPEC@@\%, AUPRC @@HIGH_AUPRC@@, and AUROC @@HIGH_AUROC@@ at @@HIGH_PREV@@\% prevalence. A patient-cross-fitted operating threshold increased sensitivity to @@HIGH_XSENS@@\% at approximately 90\% specificity, changing the task to calibrated screening rather than accurate absolute BP. For low $MAP_a$, direct sensitivity/specificity was @@LOW_SENS@@/@@LOW_SPEC@@\%, AUPRC @@LOW_AUPRC@@, and AUROC @@LOW_AUROC@@; thresholded sensitivity was @@LOW_XSENS@@\%. Errors increased sharply at pressure extremes (Fig.~\ref{fig:range}). A secondary residual interval audit on the repeated-confirmation predictions achieved @@CONF_SC@@/@@CONF_DC@@\% SBP/DBP coverage but required mean widths @@CONF_SW@@/@@CONF_DW@@ mmHg, which is too broad to imply clinical certainty.

\begin{figure}[t]\centering\includegraphics[width=\columnwidth]{../results/figures/FIGURE_3_RANGE_SPECIFIC_ERROR.pdf}\caption{Pressure-range MAE. Range imbalance makes aggregate MAE optimistic for extremes.}\label{fig:range}\end{figure}

\subsection{External transfer}
VitalDB zero-shot MAE was @@VZ_SBP@@/@@VZ_DBP@@ mmHg with bias @@VZ_BSBP@@/@@VZ_BDBP@@, versus @@VM_SBP@@/@@VM_DBP@@ for the source-training mean. Direct high-BP sensitivity was 0\%. Scores retained limited ordering information (AUPRC @@V_AUPRC@@; AUROC @@V_AUROC@@); a threshold fitted only on other destination patients gave @@V_XSENS@@\% sensitivity at @@V_XSPEC@@\% specificity. Supervised affine target-domain adaptation reached @@VA_SBP@@/@@VA_DBP@@, but this cannot support a zero-shot claim. Bland--Altman plots show wide limits and systematic external underestimation (Fig.~\ref{fig:ba}). PaPaGei was not treated as independent on VitalDB because its reported pretraining corpus included VitalDB.

\begin{figure*}[t]\centering\includegraphics[width=.92\textwidth]{../results/figures/FIGURE_1_BLAND_ALTMAN.pdf}\caption{Bland--Altman analysis for patient-nested MIMIC-BP predictions and zero-shot VitalDB predictions.}\label{fig:ba}\end{figure*}

\begin{table}[t]\caption{VitalDB patient-macro MAE (mmHg)}\centering\scriptsize
\begin{tabular}{lrr}\toprule Protocol & SBP & DBP\\\midrule Source-training mean, zero-shot & @@VM_SBP@@ & @@VM_DBP@@\\
Model, zero-shot & @@VZ_SBP@@ & @@VZ_DBP@@\\ Model, supervised affine cross-fit & @@VA_SBP@@ & @@VA_DBP@@\\\bottomrule\end{tabular}\end{table}

\section{Discussion}
The model learned population-level BP information beyond a fold-fitted mean, but the effect was modest and did not translate into safe extreme-pressure detection. This is the central biomedical result: a respectable mean error coexisted with regression toward the cohort center, poor direct high-BP sensitivity, wide limits of agreement, and external failure. Architecture novelty did not repair limited absolute-BP identifiability from normalized PPG.

Nested selection corrects the largest evaluation weakness of the earlier manuscript. It does not erase all researcher degrees of freedom because the search space was historically designed on this cohort; accordingly we distinguish the primary nested estimate from the secondary 75-fit descriptive audit. The complete registry, split replay, training-only normalization checks, and exact-once OOF coverage make that distinction auditable.

External results should be interpreted as combined domain shift. VitalDB retains simultaneous invasive numeric BP tracking but changes device, population, clinical setting, and preprocessing. Cuff datasets additionally change reference timing. Supervised cross-fitted adaptation can learn destination priors and output scaling; an improved adapted MAE therefore does not demonstrate subject-specific physiology or calibration-free deployment. Particularly large ambulatory DBP shifts and post-adaptation slopes are reported in the supplement rather than hidden.

Limitations include retrospective data, one internal cohort, outcome-informed historical design of the registry, a deterministic rather than prospective VitalDB sample, non-simultaneous cuff labels in two external cohorts, and absence of raw ABP locally for direct waveform MAP. The derived MAP event and 140/90 threshold are operational. Cross-fitted residual intervals and ensemble dispersion are exploratory. Results do not establish compliance with AAMI/ISO/ESH device-validation standards and the model must not be used for diagnosis.

\section{Conclusion}
A patient-nested 200-method selection procedure improved over a training-fold mean on MIMIC-BP but remained inaccurate at pressure extremes and @@V_CONCLUSION@@ the official source-training mean during zero-shot VitalDB transfer. Supervised target-domain adaptation reduced external error while changing the deployment claim. Patient separation, simple baselines, range bias, Bland--Altman limits, event sensitivity, and explicit calibration status are necessary alongside aggregate MAE.

\begin{thebibliography}{00}\scriptsize
\bibitem{gonzalez2023} S. Gonz\'alez, W.-T. Hsieh, and T. P.-C. Chen, ``A benchmark for machine-learning based non-invasive blood pressure estimation using photoplethysmogram,'' \emph{Scientific Data}, vol. 10, 149, 2023.
\bibitem{cisnal2026} A. Cisnal \emph{et al.}, ``Towards trustworthy AI-driven cuffless blood pressure monitoring,'' \emph{npj Digital Medicine}, 2026.
\bibitem{sanches2024} I. Sanches \emph{et al.}, ``MIMIC-BP: A curated dataset for blood pressure estimation,'' \emph{Scientific Data}, vol. 11, 1233, 2024.
\bibitem{moulaeifard2025} M. Moulaeifard, P. H. Charlton, and N. Strodthoff, ``Generalizable deep learning for photoplethysmography-based blood pressure estimation---A benchmarking study,'' \emph{Machine Learning: Health}, vol. 1, 2025.
\bibitem{pillai2025} A. Pillai \emph{et al.}, ``PaPaGei: Open foundation models for optical physiological signals,'' \emph{Proc. ICLR}, 2025.
\bibitem{liang2018} Y. Liang \emph{et al.}, ``A new, short-recorded photoplethysmogram dataset for blood pressure monitoring in China,'' \emph{Scientific Data}, vol. 5, 180020, 2018.
\bibitem{vasquez2024} S. Vasquez Salazar \emph{et al.}, ``PPG-based BP assessment dataset,'' 2024.
\bibitem{arguello2025} E. J. Arg\"uello-Prada and C. D. Casta\~no Mosquera, ``Exploring supervised machine learning models to estimate blood pressure using non-fiducial features,'' \emph{Phys. Eng. Sci. Med.}, 2025.
\bibitem{lee2022} H. C. Lee \emph{et al.}, ``VitalDB, a high-fidelity multi-parameter vital signs database in surgical patients,'' \emph{Scientific Data}, vol. 9, 279, 2022.
\end{thebibliography}
\end{document}
'''
    tex = replace(template, values)
    tex_path = PAPER / "main_reviewer_corrected.tex"
    tex_path.write_text(tex, encoding="utf-8")

    supplement = r'''\documentclass[conference]{IEEEtran}
\usepackage{booktabs,graphicx,url}\begin{document}
\title{Supplement: Reproducibility and Safety Details for Patient-Nested PPG-to-BP Estimation}\author{Anonymous submission}\maketitle
\section{Reproducibility scope}
The official 1,100/195/229 subject lists are preserved verbatim. A historical custom 1,066/229/229 split is retained only as an audited legacy artifact and is not described as official. Exact duplicate removal excludes 31 records. The input is samples 937:2812 (Python end-exclusive) from each 3,750-sample segment. All target standardization and engineered-feature transforms are fitted on the relevant training indices. The nested run stores every patient list, configuration, epoch, prediction index, checkpoint mean/SD, and journal event.
\section{Search-space disclosure}
The machine-readable table \texttt{SUPPLEMENT\_ALL\_200\_CONFIGS\_AND\_RESULTS.csv} contains every configuration and screening result; \texttt{SUPPLEMENT\_NESTED\_1005\_INNER\_FITS.csv} contains every inner fit. These are literature-informed systematic comparisons rather than causal one-factor ablations.
\section{Additional results}
Primary nested MAE was @@NEST_SBP@@/@@NEST_DBP@@ mmHg. The secondary repeated grouped audit was @@REP_SBP@@/@@REP_DBP@@. VitalDB zero-shot MAE was @@VZ_SBP@@/@@VZ_DBP@@ and supervised patient-cross-fitted affine adaptation was @@VA_SBP@@/@@VA_DBP@@. Full bias, error SD, limits of agreement, calibration slope/intercept, correlations, and within-5/10/15 rates are in \texttt{TABLE\_7\_BIAS\_BLAND\_ALTMAN\_CALIBRATION.csv}.
\begin{figure*}[t]\centering\includegraphics[width=.9\textwidth]{../results/figures/FIGURE_2_VITALDB_CALIBRATION.pdf}\caption{VitalDB reference versus model output before and after supervised target-domain adaptation.}\end{figure*}
\begin{figure*}[t]\centering\includegraphics[width=.9\textwidth]{../results/figures/FIGURE_4_SEARCH_AND_NESTED_SELECTION.pdf}\caption{Distribution of all 200 screening results and fold-specific nested selections.}\end{figure*}
\begin{figure*}[t]\centering\includegraphics[width=.9\textwidth]{../results/figures/FIGURE_5_AMBULATORY_DBP_SHIFT.pdf}\caption{Reference, zero-shot, and supervised-adapted DBP distributions in the ambulatory cohort. The large zero-shot DBP error reflects an output-scale/domain mismatch rather than clinically acceptable estimation.}\end{figure*}
\section{Claim boundaries}
No prospective clinical validation, device-standard compliance, universally diagnostic threshold, measured MAP, formal conformal guarantee, or calibrated ensemble uncertainty is claimed. VitalDB combines several domain changes. PaPaGei is excluded from independent VitalDB claims because of pretraining-corpus overlap.
\end{document}'''
    (PAPER / "supplement_reviewer_corrected.tex").write_text(replace(supplement, values), encoding="utf-8")

    tectonic = shutil.which("tectonic")
    if tectonic is None:
        raise RuntimeError("Tectonic was not found on PATH; install it before running the report stage")
    for name in ("main_reviewer_corrected.tex", "supplement_reviewer_corrected.tex"):
        process = subprocess.run([tectonic, "--outdir", str(PAPER), name], cwd=PAPER, text=True, capture_output=True)
        (PAPER / f"{Path(name).stem}_build.log").write_text(process.stdout + process.stderr, encoding="utf-8")
        if process.returncode: raise RuntimeError(f"Tectonic failed for {name}: {process.stderr[-2000:]}")

    document = Document(); title = document.add_heading("Patient-Nested Selection and Cross-Domain Safety Audit of PPG-Only Blood-Pressure Estimation", 0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    document.add_paragraph("Anonymous submission — reviewer-corrected editable draft").alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_complete_tex_narrative(document, tex)
    document.add_heading("Primary result table", level=1)
    main_table = pd.DataFrame([
        {"Model": "Nested selected procedure", "SBP MAE": selected["mae_sbp"], "DBP MAE": selected["mae_dbp"]},
        {"Model": "Historical fixed prior", "SBP MAE": prior_nested["mae_sbp"], "DBP MAE": prior_nested["mae_dbp"]},
        {"Model": "Outer-training mean", "SBP MAE": mean_nested["mae_sbp"], "DBP MAE": mean_nested["mae_dbp"]},
    ])
    add_docx_table(document, main_table, list(main_table.columns))
    for image in ("FIGURE_1_BLAND_ALTMAN.png", "FIGURE_3_RANGE_SPECIFIC_ERROR.png"):
        document.add_picture(str(ROOT / "results/figures" / image), width=Inches(6.4))
    document.add_paragraph("The complete technical manuscript and supplement are the accompanying PDF/TeX files; this DOCX preserves the editable narrative, primary table, and figures.")
    document.save(PAPER / "main_reviewer_corrected_editable.docx")
    manifest = {"paper": "paper/main_reviewer_corrected.pdf", "source": "paper/main_reviewer_corrected.tex",
                "editable": "paper/main_reviewer_corrected_editable.docx", "supplement": "paper/supplement_reviewer_corrected.pdf",
                "primary_result": selected, "review_corrections": 44}
    (PAPER / "PAPER_BUILD_MANIFEST.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    claim_registry = {
        "schema_version": "2.0", "manuscript": "paper/main_reviewer_corrected.tex",
        "note": "Exact machine values underlying the rounded manuscript statements.",
        "claims": [
            {"id": "official_split", "source": "data/mimic_bp/splits/official_patient_split.json", "value": [1100, 195, 229]},
            {"id": "retained_segments", "source": "data/mimic_bp/splits/official_patient_split.json", "value": split["counts"]["retained_segments"]},
            {"id": "excluded_duplicates", "source": "data/mimic_bp/splits/official_patient_split.json", "value": len(split["excluded_segments"])},
            {"id": "nested_inner_fits", "source": "phase2/outputs/nested200/DECISION.json", "value": nested["inner_fits"]},
            {"id": "nested_selected_metrics", "source": "phase2/outputs/nested200/DECISION.json", "value": selected},
            {"id": "nested_prior_metrics", "source": "phase2/outputs/nested200/DECISION.json", "value": prior_nested},
            {"id": "nested_mean_metrics", "source": "phase2/outputs/nested200/DECISION.json", "value": mean_nested},
            {"id": "secondary_repeated_metrics", "source": "phase2/outputs/innovation200_confirmation/DECISION.json", "value": repeated},
            {"id": "event_safety", "source": "phase2/outputs/nested200/DECISION.json", "value": event},
            {"id": "vitaldb_external", "source": "phase2/outputs/innovation200_vitaldb_external/DECISION.json", "value": vital},
            {"id": "conformal_interval_audit", "source": "phase2/outputs/icbme_clinical_analysis/CONFORMAL_90.csv", "value": conf.to_dict()},
        ],
    }
    (ROOT / "claims/final_claims.json").write_text(json.dumps(claim_registry, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(manifest, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
